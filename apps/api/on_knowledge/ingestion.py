from __future__ import annotations

import hashlib
import http.client
import ipaddress
import re
import socket
import ssl
import zipfile
from pathlib import Path
from urllib.parse import urljoin, urlsplit

from bs4 import BeautifulSoup
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, Field
from pypdf import PdfReader

MAX_BYTES = 25 * 1024 * 1024
MAX_TEXT = 8_000_000
EXTENSIONS = {".pdf", ".md", ".txt", ".docx"}


class Element(BaseModel):
    type: str = "paragraph"
    text: str
    page: int | None = None
    section: str = ""
    line_start: int | None = None
    line_end: int | None = None


class NormalizedDocument(BaseModel):
    elements: list[Element] = Field(default_factory=list)
    parser: str


def text_elements(text: str, page: int | None = None) -> list[Element]:
    result, lines, heading, start, code = [], [], "", 1, False

    def flush(end):
        if lines:
            result.append(Element(text="\n".join(lines), page=page, section=heading,
                                  type="code" if code else "paragraph", line_start=start, line_end=end))
            lines.clear()

    for i, line in enumerate(text.splitlines(), 1):
        if line.startswith("```"):
            if code:
                lines.append(line)
                flush(i)
                code = False
            else:
                flush(i - 1)
                code, start = True, i
                lines.append(line)
        elif not code and re.match(r"^#{1,6}\s", line):
            flush(i - 1)
            heading = line.lstrip("# ")
            result.append(Element(type="heading", text=line, page=page, section=heading,
                                  line_start=i, line_end=i))
        elif not code and not line.strip():
            flush(i - 1)
        else:
            if not lines:
                start = i
            lines.append(line)
    flush(len(text.splitlines()))
    return result


def parse(path: Path) -> NormalizedDocument:
    suffix = path.suffix.lower()
    if suffix not in EXTENSIONS:
        raise ValueError("Поддерживаются PDF, Markdown, TXT и DOCX")
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("Размер файла превышает 25 МБ")
    if suffix == ".pdf":
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Сначала снимите пароль с PDF")
        if len(reader.pages) > 2000:
            raise ValueError("PDF превышает 2000 страниц")
        elements = []
        total = 0
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text(extraction_mode="layout")
            total += len(text)
            if total > MAX_TEXT:
                raise ValueError("Слишком большой извлечённый текст")
            elements.extend(text_elements(text, page=number))
        parser = "pypdf-layout-v1"
    elif suffix == ".docx":
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if (len(entries) > 3000 or sum(f.file_size for f in entries) > 100 * 1024 * 1024
                    or any(f.file_size > 1000 * max(f.compress_size, 1) for f in entries)):
                raise ValueError("DOCX превышает лимит распаковки")
        elements, section = [], ""
        for block in Document(path).iter_inner_content():
            if isinstance(block, Paragraph) and block.text.strip():
                is_heading = block.style and block.style.name.startswith("Heading")
                if is_heading:
                    section = block.text
                elements.append(Element(type="heading" if is_heading else "paragraph",
                                        text=block.text, section=section))
            elif isinstance(block, Table):
                elements.append(Element(type="table", section=section,
                                        text="\n".join(" | ".join(c.text for c in row.cells)
                                                       for row in block.rows)))
        parser = "python-docx-structure-v1"
    else:
        elements, parser = text_elements(path.read_text("utf-8-sig")), "markdown-blocks-v1"
    if not elements:
        raise ValueError("Текст не найден. Для сканированного PDF требуется OCR (ещё не включён).")
    if sum(len(e.text) for e in elements) > MAX_TEXT:
        raise ValueError("Слишком большой извлечённый текст")
    return NormalizedDocument(elements=elements, parser=parser)


def chunk_document(document: NormalizedDocument, source: str, version: str) -> list[dict]:
    """Preserve element boundaries and exact quote offsets; long blocks split on words."""
    chunks = []
    for index, element in enumerate(document.elements):
        matches = list(re.finditer(r"\S+", element.text))
        for start in range(0, len(matches), 300):
            group = matches[start:start + 340]
            a, b = group[0].start(), group[-1].end()
            text = element.text[a:b]
            location = {**element.model_dump(exclude={"text"}), "element": index,
                        "offset_start": a, "offset_end": b}
            key = hashlib.sha256(f"{source}:{version}:{index}:{a}:{b}".encode()).hexdigest()
            chunks.append({"id": key, "text": text, "location": location})
            if start + 340 >= len(matches):
                break
    return chunks


def public_addresses(host: str, port: int) -> list[str]:
    addresses = list({info[4][0] for info in socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)})
    if not addresses or any(not ipaddress.ip_address(ip).is_global for ip in addresses):
        raise ValueError("URL должен вести на публичный сайт; локальные адреса запрещены")
    return addresses


def fetch_url(url: str) -> tuple[str, str]:
    """Pin a validated DNS address for the actual socket, including redirects (no rebinding)."""
    for _ in range(5):
        parsed = urlsplit(url)
        if parsed.scheme not in ("http", "https") or not parsed.hostname or parsed.username or parsed.password:
            raise ValueError("Допустим только публичный HTTP(S) URL без пароля")
        port = parsed.port or (443 if parsed.scheme == "https" else 80)
        if port not in (80, 443):
            raise ValueError("URL должен использовать порт 80 или 443")
        addresses = public_addresses(parsed.hostname, port)
        connection = http.client.HTTPConnection(parsed.hostname, port, timeout=20)
        sock = socket.create_connection((addresses[0], port), timeout=20)
        if parsed.scheme == "https":
            sock = ssl.create_default_context().wrap_socket(sock, server_hostname=parsed.hostname)
        connection.sock = sock
        try:
            connection.request("GET", (parsed.path or "/") + ("?" + parsed.query if parsed.query else ""),
                               headers={"User-Agent": "OpenNotebook/0.1", "Accept-Encoding": "identity"})
            response = connection.getresponse()
            if response.status in (301, 302, 303, 307, 308):
                url = urljoin(url, response.getheader("Location", ""))
                continue
            if response.status != 200:
                raise ValueError(f"Сайт вернул HTTP {response.status}")
            content_type = response.getheader("Content-Type", "").lower()
            if not any(t in content_type for t in ("text/html", "text/plain", "application/xhtml")):
                raise ValueError("URL должен возвращать HTML или текст")
            data = response.read(MAX_BYTES + 1)
            if len(data) > MAX_BYTES:
                raise ValueError("Страница превышает 25 МБ")
            soup = BeautifulSoup(data, "html.parser")
            for tag in soup(["script", "style", "nav", "footer", "iframe", "form"]):
                tag.decompose()
            title = soup.title.get_text(strip=True) if soup.title else parsed.hostname
            for tag in soup.find_all(re.compile(r"^h[1-6]$")):
                tag.replace_with("\n" + "#" * int(tag.name[1]) + " " + tag.get_text(" ", strip=True) + "\n")
            return title, soup.get_text("\n", strip=True)
        finally:
            connection.close()
    raise ValueError("Слишком много перенаправлений")
