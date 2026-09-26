import io
import socket
import time
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from on_knowledge.app import create_app
from on_knowledge.ingestion import Element, NormalizedDocument, chunk_document, parse, public_addresses
from on_knowledge.knowledge import cosine, rrf
from on_knowledge.providers import OpenAICompatible, validate_endpoint
from on_knowledge.storage import write_json
from reportlab.pdfgen.canvas import Canvas


@pytest.fixture
def client(tmp_path):
    app = create_app(tmp_path, "test-session-token-long-enough")
    with TestClient(app, headers={"Authorization": "Bearer test-session-token-long-enough"}) as c:
        yield c


def create(client, name="Research"):
    response = client.post("/api/notebooks", json={"title": name})
    assert response.status_code == 201
    return response.json()["id"]


def upload(client, notebook, filename="facts.md", data=b"# Project\n\nThe codename is Marigold. Launch is October 18."):
    response = client.post(f"/api/notebooks/{notebook}/sources", files={"file": (filename, data)})
    assert response.status_code == 202, response.text
    source = response.json()
    for _ in range(100):
        state = client.get(f'/api/notebooks/{notebook}/sources/{source["id"]}').json()["source"]
        if state["status"] in ("ready", "error"):
            assert state["status"] == "ready", state
            return state
        time.sleep(.03)
    pytest.fail("Ingestion did not finish")


def test_vertical_slice_pdf_provenance_and_reindex(client):
    notebook = create(client)
    pdf = io.BytesIO()
    canvas = Canvas(pdf)
    canvas.drawString(50, 750, "A preliminary page without the answer.")
    canvas.showPage()
    canvas.drawString(50, 750, "The Marigold mission reaches Europa on October 18.")
    canvas.save()
    source = upload(client, notebook, "mission.pdf", pdf.getvalue())
    evidence = client.post(f"/api/notebooks/{notebook}/search", json={"question": "Marigold Europa"}).json()["evidence"]
    assert evidence[0]["location"]["page"] == 2
    assert evidence[0]["source_id"] == source["id"]
    assert "October 18" in evidence[0]["quote"]
    document = client.get(f'/api/notebooks/{notebook}/sources/{source["id"]}').json()["document"]
    loc = evidence[0]["location"]
    assert document["elements"][loc["element"]]["text"][loc["offset_start"]:loc["offset_end"]] == evidence[0]["quote"]
    original = client.get(f'/api/notebooks/{notebook}/sources/{source["id"]}/original').content
    assert original == pdf.getvalue()
    knowledge = client.app.state.knowledge
    knowledge.ingest(notebook, source["id"])
    rebuilt = knowledge.search(notebook, "Marigold Europa")["evidence"]
    assert evidence[0]["id"] == rebuilt[0]["id"]


def test_auth_and_notebook_isolation(client):
    a, b = create(client, "A"), create(client, "B")
    source = upload(client, a)
    assert client.get("/api/notebooks", headers={"Authorization": "Bearer wrong"}).status_code == 401
    assert client.post(f"/api/notebooks/{b}/search", json={"question": "Marigold"}).json()["evidence"] == []
    assert client.get(f'/api/notebooks/{b}/sources/{source["id"]}').status_code == 404


def test_notes_frontmatter_and_external_edit(client):
    notebook = create(client)
    note = client.post(f"/api/notebooks/{notebook}/notes", json={"title": "A: title", "body": "## Idea\n\n[[Other note]]"}).json()
    path = client.app.state.store.notebook_path(notebook) / "notes" / (note["id"] + ".md")
    assert "[[Other note]]" in path.read_text("utf-8")
    path.write_text(path.read_text("utf-8") + "\nExternal edit", "utf-8")
    notes = client.get(f"/api/notebooks/{notebook}/notes").json()
    assert notes[0]["title"] == "A: title"
    assert "External edit" in notes[0]["body"]
    traversal = client.post(f"/api/notebooks/{notebook}/notes", json={"id": "../../outside", "title": "Escape", "body": "x"})
    assert traversal.status_code == 400


def test_unsupported_and_empty_upload(client):
    notebook = create(client)
    for name, content in [("file.exe", b"MZ"), ("empty.md", b"")]:
        assert client.post(f"/api/notebooks/{notebook}/sources", files={"file": (name, content)}).status_code == 400


@pytest.mark.parametrize("ip", ["127.0.0.1", "10.0.0.1", "169.254.169.254", "::1", "192.168.2.1", "0.0.0.0"])
def test_ssrf_blocks_private_resolutions(monkeypatch, ip):
    monkeypatch.setattr(socket, "getaddrinfo", lambda *args, **kwargs: [(2, 1, 6, "", (ip, 443))])
    with pytest.raises(ValueError):
        public_addresses("attacker.example", 443)


def test_provider_url_guard():
    assert validate_endpoint("http://127.0.0.1:11434/v1/") == "http://127.0.0.1:11434/v1"
    for url in ["http://example.com/v1", "https://user:password@example.com", "file:///etc/passwd"]:
        with pytest.raises(ValueError):
            validate_endpoint(url)


def test_docx_structure_and_archive_bomb(tmp_path):
    doc = Document()
    doc.add_heading("Decisions", 1)
    doc.add_paragraph("Use a local knowledge store.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Retention", "365 days"
    path = tmp_path / "file.docx"
    doc.save(path)
    parsed = parse(path)
    assert parsed.elements[1].section == "Decisions"
    assert parsed.elements[2].type == "table"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("bomb", b"0" * 2_000_000)
    with pytest.raises(ValueError, match="распаковки"):
        parse(path)


def test_chunk_offsets_no_truncation():
    text = " ".join(f"word{i}" for i in range(1031))
    doc = NormalizedDocument(parser="fixture", elements=[Element(type="code", text=text, page=4, section="Code")])
    chunks = chunk_document(doc, "s", "v")
    for chunk in chunks:
        loc = chunk["location"]
        assert chunk["text"] == text[loc["offset_start"]:loc["offset_end"]]
        assert loc["page"] == 4 and loc["type"] == "code"
    assert "word1030" in chunks[-1]["text"]
    assert all(f"word{i}" in " ".join(c["text"] for c in chunks) for i in range(1031))


def test_rrf_and_cosine():
    assert rrf([["a", "b"], ["b", "c"]])[0][0] == "b"
    assert rrf([["a", "a"]])[0][1] == 1 / 61
    assert cosine([1, 0], [1, 0]) == 1
    assert cosine([0, 0], [1, 0]) == 0


def test_answer_rejects_invented_citations_and_critic_failures(client, monkeypatch):
    notebook = create(client)
    upload(client, notebook)
    def complete(self, system, data):
        if "supported_indices" in system:
            return {"supported_indices": [0]}
        return {"claims": [
            {"text": "The codename is Marigold.", "evidence_ids": [data["evidence"][0]["id"]]},
            {"text": "Invented!", "evidence_ids": ["nonexistent"]},
            {"text": "Unsupported", "evidence_ids": [data["evidence"][0]["id"]]},
        ]}
    monkeypatch.setattr(OpenAICompatible, "complete", complete)
    response = client.post(f"/api/notebooks/{notebook}/ask", json={"question": "Marigold"})
    assert response.status_code == 200
    answer = response.json()
    assert len(answer["claims"]) == 1
    assert answer["citations"][0]["number"] == 1
    assert answer["validated"]


def test_research_and_memory_gate(client, monkeypatch):
    notebook = create(client)
    upload(client, notebook)
    def complete(self, system, data):
        if "Plan research" in system:
            return {"queries": ["codename"]}
        if "supported_indices" in system:
            return {"supported_indices": [0]}
        return {"claims": [{"text": "Marigold is the codename.", "evidence_ids": [data["evidence"][0]["id"]]}]}
    monkeypatch.setattr(OpenAICompatible, "complete", complete)
    session = client.post(f"/api/notebooks/{notebook}/research", json={"question": "Marigold"}).json()
    assert session["status"] == "completed"
    assert client.post(f'/api/notebooks/{notebook}/research/{session["id"]}/save').status_code == 200
    root = client.app.state.store.notebook_path(notebook)
    assert (root / "memory" / (session["id"] + ".json")).exists()
    session["answer"]["validated"] = False
    write_json(root / "research" / (session["id"] + ".json"), session)
    assert client.post(f'/api/notebooks/{notebook}/research/{session["id"]}/save').status_code == 400


def test_rebuild_after_index_loss(client):
    notebook = create(client)
    source = upload(client, notebook)
    store = client.app.state.store
    with store.db() as db:
        db.execute("DELETE FROM chunks")
        db.execute("DELETE FROM chunks_fts")
    client.app.state.knowledge.recover()
    for _ in range(100):
        if client.app.state.knowledge.search(notebook, "Marigold")["evidence"]:
            break
        time.sleep(.03)
    assert client.app.state.knowledge.search(notebook, "Marigold")["evidence"]
    assert (store.source_path(notebook, source["id"]) / source["original"]).exists()


def test_cors_preflight_and_frontmatter_preservation(client):
    response = client.options("/api/notebooks", headers={
        "Origin": "http://localhost:3000", "Access-Control-Request-Method": "POST",
        "Access-Control-Request-Headers": "authorization,content-type",
        "Authorization": "",
    })
    assert response.status_code == 200
    notebook = create(client)
    note = client.post(f"/api/notebooks/{notebook}/notes", json={"title": "Human title", "body": "Before"}).json()
    assert note["id"] == "Human title"
    path = client.app.state.store.notebook_path(notebook) / "notes" / "Human title.md"
    path.write_text("---\ntitle: Human title\ntags: [retained]\ncustom: value\n---\n\nBefore", "utf-8")
    # Refresh after editing externally: a stale revision must no longer overwrite the file.
    note = client.get(f"/api/notebooks/{notebook}/notes").json()[0]
    assert client.post(f"/api/notebooks/{notebook}/notes", json={**note, "body": "After"}).status_code == 200
    assert "retained" in path.read_text("utf-8") and "custom: value" in path.read_text("utf-8")


def test_hybrid_retrieves_without_keyword_overlap_and_rejects_model_mismatch(client, monkeypatch):
    notebook = create(client)
    write_json(client.app.state.store.root / "settings.json", {
        "base_url": "http://127.0.0.1:11434/v1", "model": "", "embedding_model": "fixture-embedding"
    })
    monkeypatch.setattr(OpenAICompatible, "embed", lambda self, texts: [[1., 0.] for _ in texts])
    upload(client, notebook)
    result = client.post(f"/api/notebooks/{notebook}/search", json={"question": "unrelated lexical terms"}).json()
    assert result["mode"] == "hybrid"
    assert result["evidence"]
    write_json(client.app.state.store.root / "settings.json", {
        "base_url": "http://127.0.0.1:11434/v1", "model": "", "embedding_model": "different-model"
    })
    result = client.post(f"/api/notebooks/{notebook}/search", json={"question": "unrelated lexical terms"}).json()
    assert result["evidence"] == []
